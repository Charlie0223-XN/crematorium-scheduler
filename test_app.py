"""Flask API 與 Excel v2 整合測試。執行：python -m unittest -v test_app.py"""
from __future__ import annotations

import io
import unittest
from datetime import date, timedelta

from openpyxl import load_workbook

from app import app
from scheduler import EMPLOYEES


START = date(2026, 9, 1)


def valid_payload():
    days = [
        {
            "date": (START + timedelta(days=index)).isoformat(),
            "day_type": "NORMAL",
            "label": "",
            "requirements": {},
        }
        for index in range(28)
    ]
    days[4]["day_type"] = "BIG"
    days[9]["day_type"] = "OFF"
    days[18] = {
        "date": (START + timedelta(days=18)).isoformat(),
        "day_type": "CUSTOM",
        "label": "停爐留守",
        "requirements": {"A": 0, "B": 2, "C": 0},
    }

    vacations = {}
    for employee_index, name in enumerate(EMPLOYEES):
        indexes = []
        candidate = employee_index
        while len(indexes) < 8:
            index = candidate % 28
            if index != 9 and index not in indexes:
                indexes.append(index)
            candidate += 13
        vacations[name] = [
            (START + timedelta(days=index)).isoformat() for index in sorted(indexes)
        ]

    return {"seed": 123456, "days": days, "vacations": vacations}


class AppV2Tests(unittest.TestCase):
    def setUp(self):
        app.config.update(TESTING=True)
        self.client = app.test_client()

    def test_home_and_health(self):
        home = self.client.get("/")
        self.assertEqual(home.status_code, 200)
        self.assertIn("新廠四週排班".encode("utf-8"), home.data)

        health = self.client.get("/api/health")
        self.assertEqual(health.status_code, 200)
        self.assertEqual(health.get_json(), {"ok": True, "version": 2})

    def test_generate_valid_schedule(self):
        response = self.client.post("/api/generate_final", json=valid_payload())
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        data = response.get_json()
        self.assertEqual(len(data["schedule"]), 28)
        self.assertEqual(data["schedule"][9]["assignment"], {})
        self.assertEqual(len(data["schedule"][18]["assignment"]), 2)
        self.assertEqual(len(data["stats"]["employees"]), 13)

    def test_requires_28_consecutive_days(self):
        payload = valid_payload()
        payload["days"].pop()
        response = self.client.post("/api/generate_final", json=payload)
        self.assertEqual(response.status_code, 400)
        self.assertIn("28", response.get_json()["error"])

    def test_requires_eight_vacation_days_per_employee(self):
        payload = valid_payload()
        payload["vacations"]["子紘"].pop()
        response = self.client.post("/api/generate_final", json=payload)
        self.assertEqual(response.status_code, 400)
        self.assertIn("子紘", response.get_json()["error"])

    def test_custom_requirement_cannot_exceed_available_people(self):
        payload = valid_payload()
        payload["days"][18]["requirements"] = {"A": 8, "B": 8, "C": 8}
        response = self.client.post("/api/generate_final", json=payload)
        self.assertEqual(response.status_code, 400)
        self.assertIn("需求", response.get_json()["error"])

    def test_export_contains_three_expected_sheets(self):
        response = self.client.post("/api/export_excel", json=valid_payload())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(io.BytesIO(response.data), read_only=True)
        self.assertEqual(workbook.sheetnames, ["每日班表", "人員統計", "休假設定"])
        self.assertEqual(workbook["每日班表"].max_row, 29)
        self.assertEqual(workbook["人員統計"].max_row, 16)


if __name__ == "__main__":
    unittest.main(verbosity=2)
